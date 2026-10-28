from __future__ import annotations

import hashlib
import zipfile
from pathlib import Path

import openpyxl
import pymupdf as fitz

from .config import Config
from .models import ExtractedDocument, QueuedJob
from .normalizer import normalize

MAX_ZIP_ENTRIES = 2_000
MAX_UNCOMPRESSED_BYTES = 80 * 1024 * 1024


def resolve_document(storage_root: Path, job: QueuedJob) -> Path:
    candidate = (storage_root / job.document.storageId).resolve(strict=True)
    if (
        candidate.parent != storage_root.resolve()
        or candidate.is_symlink()
        or not candidate.is_file()
    ):
        raise ValueError("unsafe document path")
    stat = candidate.stat()
    if stat.st_size != job.document.size:
        raise ValueError("document size mismatch")
    if _sha256(candidate) != job.document.sha256:
        raise ValueError("document digest mismatch")
    return candidate


def extract_document(path: Path, job: QueuedJob, config: Config) -> ExtractedDocument:
    if path.suffix == ".pdf":
        text, warnings = _extract_pdf(path, config)
        result = ExtractedDocument(
            job_id=job.jobId,
            request_id=job.requestId,
            kind="pdf",
            text=text,
            warnings=warnings,
        )
        result.normalized = normalize(result.text)
        return result
    if path.suffix == ".xlsx":
        text, sheets, warnings = _extract_xlsx(path, config)
        result = ExtractedDocument(
            job_id=job.jobId,
            request_id=job.requestId,
            kind="xlsx",
            text=text,
            sheets=sheets,
            warnings=warnings,
        )
        result.normalized = normalize(result.text)
        return result
    raise ValueError("unsupported document type")


def _extract_pdf(path: Path, config: Config) -> tuple[str, list[str]]:
    with path.open("rb") as stream:
        if stream.read(5) != b"%PDF-":
            raise ValueError("invalid PDF signature")
    warnings: list[str] = []
    fragments: list[str] = []
    count = 0
    with fitz.open(path) as document:  # type: ignore[no-untyped-call]
        if document.page_count > config.max_pdf_pages:
            warnings.append(f"processed first {config.max_pdf_pages} pages")
        for page_number in range(min(document.page_count, config.max_pdf_pages)):
            fragment = document.load_page(page_number).get_text("text")
            remaining = config.max_extracted_chars - count
            if remaining <= 0:
                warnings.append("text truncated")
                break
            fragments.append(fragment[:remaining])
            count += len(fragments[-1])
    return _normalise_text("\n".join(fragments)), warnings


def _extract_xlsx(path: Path, config: Config) -> tuple[str, list[str], list[str]]:
    _validate_zip(path)
    warnings: list[str] = []
    fragments: list[str] = []
    sheet_names: list[str] = []
    count = 0
    workbook = openpyxl.load_workbook(path, read_only=True, data_only=True, keep_links=False)
    try:
        for sheet in workbook.worksheets:
            sheet_names.append(sheet.title[:100])
            fragments.append(f"[Лист: {sheet.title[:100]}]")
            for row in sheet.iter_rows(values_only=True):
                line = "\t".join(_safe_cell(value) for value in row if value is not None)
                if not line:
                    continue
                remaining = config.max_extracted_chars - count
                if remaining <= 0:
                    warnings.append("text truncated")
                    return _normalise_text("\n".join(fragments)), sheet_names, warnings
                fragments.append(line[:remaining])
                count += len(fragments[-1])
    finally:
        workbook.close()
    return _normalise_text("\n".join(fragments)), sheet_names, warnings


def _validate_zip(path: Path) -> None:
    total = 0
    with zipfile.ZipFile(path) as archive:
        names = set(archive.namelist())
        if "[Content_Types].xml" not in names or "xl/workbook.xml" not in names:
            raise ValueError("invalid XLSX structure")
        if len(names) > MAX_ZIP_ENTRIES:
            raise ValueError("too many XLSX entries")
        for entry in archive.infolist():
            parts = Path(entry.filename).parts
            if entry.filename.startswith(("/", "\\")) or ".." in parts or "\\" in entry.filename:
                raise ValueError("unsafe XLSX entry")
            total += entry.file_size
            if total > MAX_UNCOMPRESSED_BYTES:
                raise ValueError("XLSX expanded size exceeded")


def _sha256(path: Path) -> str:
    digest = hashlib.sha256()
    with path.open("rb") as stream:
        for chunk in iter(lambda: stream.read(1024 * 1024), b""):
            digest.update(chunk)
    return digest.hexdigest()


def _safe_cell(value: object) -> str:
    text = str(value).replace("\x00", "").replace("\r", " ").replace("\n", " ")
    # Prevent formula injection if extracted values are later exported to a spreadsheet.
    if text.lstrip().startswith(("=", "+", "-", "@")):
        text = "'" + text
    return text[:10_000]


def _normalise_text(text: str) -> str:
    return "\n".join(line.rstrip() for line in text.replace("\x00", "").splitlines()).strip()

